# -*- coding: utf-8 -*-
"""
Generator raportu dla klienta (.docx) z weryfikacji integracji IdoSell ↔ Meta Ads.

Użycie:
    python raport_docx.py <folder_raportu>

Folder raportu (np. raporty/pyrawinyl.pl_2026-10-14) musi zawierać dane.json oraz zrzuty/.
Wynik: <folder>/Raport_Meta_<sklep>_<data>.docx oraz <folder>/komentarz_clickup.txt.
"""

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

STATUSY = {
    "ok": ("W porządku", "2E7D32"),
    "uwaga": ("Do poprawy", "EF6C00"),
    "problem": ("Wymaga działania", "C62828"),
}
KOLOR_NAGLOWKA = RGBColor(0x1F, 0x3A, 0x5F)


def cieniuj_komorke(komorka, kolor_hex):
    tc_pr = komorka._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kolor_hex)
    tc_pr.append(shd)


def akapit(doc, tekst, rozmiar=10.5, pogrubiony=False, kolor=None, odstep_po=4):
    p = doc.add_paragraph()
    run = p.add_run(tekst)
    run.font.size = Pt(rozmiar)
    run.bold = pogrubiony
    if kolor:
        run.font.color.rgb = kolor
    p.paragraph_format.space_after = Pt(odstep_po)
    return p


def etykieta_statusu(doc, status):
    nazwa, kolor = STATUSY[status]
    p = doc.add_paragraph()
    run = p.add_run(f"● {nazwa}")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(kolor)
    p.paragraph_format.space_after = Pt(2)


def buduj(folder):
    dane = json.loads((folder / "dane.json").read_text(encoding="utf-8"))
    doc = Document()

    styl = doc.styles["Normal"]
    styl.font.name = "Calibri"
    styl.font.size = Pt(10.5)
    for sekcja in doc.sections:
        sekcja.left_margin = sekcja.right_margin = Cm(2)
        sekcja.top_margin = sekcja.bottom_margin = Cm(1.8)

    # Nagłówek
    tytul = doc.add_heading(level=0)
    run = tytul.add_run("Weryfikacja integracji z Meta Ads")
    run.font.color.rgb = KOLOR_NAGLOWKA
    akapit(doc, f"Sklep: {dane['sklep']}", 12, True, odstep_po=0)
    akapit(doc, f"Data weryfikacji: {dane['data']}    Okres analizy: {dane['okres']}", 10, odstep_po=12)

    # Podsumowanie
    doc.add_heading("Podsumowanie", level=1)
    for punkt in dane["podsumowanie"]:
        p = doc.add_paragraph(punkt, style="List Bullet")
        p.paragraph_format.space_after = Pt(2)

    # Tabela stanu obszarów
    doc.add_heading("Stan obszarów", level=1)
    tabela = doc.add_table(rows=1, cols=3)
    tabela.style = "Table Grid"
    tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
    for komorka, tekst in zip(tabela.rows[0].cells, ("Obszar", "Stan", "Wynik")):
        komorka.text = tekst
        komorka.paragraphs[0].runs[0].bold = True
        cieniuj_komorke(komorka, "DCE6F0")
    for sekcja in dane["sekcje"]:
        wiersz = tabela.add_row().cells
        nazwa, kolor = STATUSY[sekcja["status"]]
        wiersz[0].text = sekcja["tytul"]
        wiersz[1].text = nazwa
        wiersz[1].paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(kolor)
        wiersz[1].paragraphs[0].runs[0].bold = True
        wiersz[2].text = sekcja["wynik"]
    for wiersz in tabela.rows:
        wiersz.cells[0].width = Cm(5)
        wiersz.cells[1].width = Cm(3.2)
        wiersz.cells[2].width = Cm(8.8)

    # Sekcje szczegółowe
    for sekcja in dane["sekcje"]:
        doc.add_heading(sekcja["tytul"], level=1)
        etykieta_statusu(doc, sekcja["status"])
        for linia in sekcja["opis"]:
            akapit(doc, linia)
        for zrzut in sekcja.get("zrzuty", []):
            sciezka = folder / "zrzuty" / zrzut["plik"]
            if not sciezka.exists():
                akapit(doc, f"[brak zrzutu: {zrzut['plik']}]", 9, kolor=RGBColor(0x99, 0, 0))
                continue
            doc.add_picture(str(sciezka), width=Cm(16.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            podpis = akapit(doc, zrzut["podpis"], 9, kolor=RGBColor(0x55, 0x55, 0x55), odstep_po=10)
            podpis.alignment = WD_ALIGN_PARAGRAPH.CENTER
            podpis.runs[0].italic = True
        if sekcja.get("rekomendacja"):
            p = akapit(doc, "Rekomendacja: ", pogrubiony=True, odstep_po=10)
            p.add_run(sekcja["rekomendacja"]).bold = False

    # Plan działań
    doc.add_heading("Zalecane działania", level=1)
    tabela = doc.add_table(rows=1, cols=3)
    tabela.style = "Table Grid"
    for komorka, tekst in zip(tabela.rows[0].cells, ("Priorytet", "Działanie", "Po czyjej stronie")):
        komorka.text = tekst
        komorka.paragraphs[0].runs[0].bold = True
        cieniuj_komorke(komorka, "DCE6F0")
    for dzialanie in dane["dzialania"]:
        wiersz = tabela.add_row().cells
        wiersz[0].text = dzialanie["priorytet"]
        wiersz[1].text = dzialanie["dzialanie"]
        wiersz[2].text = dzialanie["kto"]
    for wiersz in tabela.rows:
        wiersz.cells[0].width = Cm(2.6)
        wiersz.cells[1].width = Cm(11)
        wiersz.cells[2].width = Cm(3.4)

    akapit(doc, "", odstep_po=6)
    akapit(doc, f"Kolejna weryfikacja: {dane['kolejna_weryfikacja']}", 10, kolor=RGBColor(0x55, 0x55, 0x55))

    nazwa = f"Raport_Meta_{dane['sklep']}_{dane['data']}.docx"
    doc.save(folder / nazwa)
    (folder / "komentarz_clickup.txt").write_text(dane["komentarz_clickup"] + "\n", encoding="utf-8")
    print(f"Zapisano: {folder / nazwa}")
    print(f"Zapisano: {folder / 'komentarz_clickup.txt'}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    buduj(Path(sys.argv[1]))
